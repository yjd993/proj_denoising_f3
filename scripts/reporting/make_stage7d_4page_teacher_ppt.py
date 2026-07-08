from __future__ import annotations

import os
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt


ROOT = Path(
    os.environ.get(
        "AIRBORNE_LIDAR_PROJECT_DATA_ROOT",
        r"E:\yjind_work\Airborne_LiDAR_Data\proj_denoising_f3",
    )
)
OUT_DIR = ROOT / "阶段汇报"
PACKAGE_DIR = (
    Path(
        os.environ.get(
            "AIRBORNE_LIDAR_STAGE_ROOT",
            r"E:\yjind_work\Airborne_LiDAR_Data\proj_denoising_f3_2.0",
        )
    )
    / "stage7d_teacher_package"
)
FIG_DIR = PACKAGE_DIR / "matlab_figures"
PPTX_PATH = OUT_DIR / "Stage7D_L1_L2_L3_处理流程_4页汇报.pptx"

FONT_CN = "Microsoft YaHei"
NAVY = RGBColor(22, 50, 78)
BLUE = RGBColor(39, 107, 164)
LIGHT_BLUE = RGBColor(229, 241, 250)
GRAY = RGBColor(88, 96, 105)
LIGHT_GRAY = RGBColor(245, 247, 249)
GREEN = RGBColor(34, 128, 93)
ORANGE = RGBColor(211, 114, 39)
RED = RGBColor(180, 68, 63)
WHITE = RGBColor(255, 255, 255)


def set_run(run, size=18, bold=False, color=RGBColor(30, 30, 30)) -> None:
    run.font.name = FONT_CN
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color


def add_textbox(slide, x, y, w, h, text, size=18, bold=False, color=RGBColor(30, 30, 30), align=None):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    p = tf.paragraphs[0]
    if align is not None:
        p.alignment = align
    run = p.add_run()
    run.text = text
    set_run(run, size=size, bold=bold, color=color)
    return box


def add_title(slide, title, subtitle=None) -> None:
    add_textbox(slide, 0.45, 0.23, 12.35, 0.45, title, size=25, bold=True, color=NAVY)
    line = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, Inches(0.45), Inches(0.82), Inches(12.45), Inches(0.03))
    line.fill.solid()
    line.fill.fore_color.rgb = BLUE
    line.line.fill.background()
    if subtitle:
        add_textbox(slide, 0.48, 0.88, 12.2, 0.28, subtitle, size=10.5, color=GRAY)


def add_note(slide, text) -> None:
    add_textbox(slide, 0.52, 7.05, 12.15, 0.22, text, size=8.5, color=GRAY)


def add_panel(slide, x, y, w, h, title=None, fill=LIGHT_GRAY, line=RGBColor(222, 226, 230)):
    shape = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.color.rgb = line
    shape.line.width = Pt(0.75)
    if title:
        add_textbox(slide, x + 0.18, y + 0.12, w - 0.36, 0.28, title, size=13, bold=True, color=NAVY)
    return shape


def add_bullets(slide, x, y, w, h, items, size=13, color=RGBColor(35, 35, 35), gap=True):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.level = 0
        p.text = item
        p.font.name = FONT_CN
        p.font.size = Pt(size)
        p.font.color.rgb = color
        p.space_after = Pt(5 if gap else 1)
    return box


def add_flow_box(slide, x, y, w, h, title, subtitle, color):
    shape = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()
    add_textbox(slide, x + 0.12, y + 0.13, w - 0.24, 0.26, title, size=13, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
    add_textbox(slide, x + 0.12, y + 0.46, w - 0.24, 0.36, subtitle, size=9.5, color=WHITE, align=PP_ALIGN.CENTER)


def add_arrow(slide, x, y, w, h):
    arrow = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RIGHT_ARROW, Inches(x), Inches(y), Inches(w), Inches(h))
    arrow.fill.solid()
    arrow.fill.fore_color.rgb = RGBColor(190, 201, 211)
    arrow.line.fill.background()


def add_picture_fit(slide, path: Path, x, y, w, h):
    if not path.exists():
        add_panel(slide, x, y, w, h, "缺少图片", fill=RGBColor(255, 245, 245), line=RED)
        add_textbox(slide, x + 0.2, y + 0.55, w - 0.4, h - 0.7, str(path), size=10, color=RED)
        return
    slide.shapes.add_picture(str(path), Inches(x), Inches(y), width=Inches(w), height=Inches(h))


def build_slide1(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_title(slide, "机载激光测深数据处理流程：从 L1 原始观测到最终点云", "以 00111 文件为例说明中间过程；00050-00150 的 95 个 Stage 7A 文件作为可展示候选成果")

    add_textbox(slide, 0.7, 1.25, 11.9, 0.35, "核心问题：原始激光记录怎样一步步变成地图坐标中的三维点云？", size=18, bold=True, color=NAVY, align=PP_ALIGN.CENTER)

    y = 2.03
    add_flow_box(slide, 0.65, y, 2.45, 0.95, "L1 原始观测", "时间 + 原始距离计数 + 扫描镜位置", BLUE)
    add_arrow(slide, 3.18, y + 0.27, 0.58, 0.38)
    add_flow_box(slide, 3.85, y, 2.45, 0.95, "L2 机体系点云", "相对飞机的前后、左右、上下位置", GREEN)
    add_arrow(slide, 6.38, y + 0.27, 0.58, 0.38)
    add_flow_box(slide, 7.05, y, 2.45, 0.95, "L3 POS 匹配", "匹配飞机位置和姿态，得到东/北/高偏移", ORANGE)
    add_arrow(slide, 9.58, y + 0.27, 0.58, 0.38)
    add_flow_box(slide, 10.25, y, 2.45, 0.95, "最终点云", "UTM 坐标 + 高程 + 经纬度", RED)

    add_panel(slide, 0.75, 3.55, 5.9, 2.45, "给不熟悉数据的老师先讲清楚")
    add_bullets(
        slide,
        1.02,
        4.02,
        5.35,
        1.65,
        [
            "L1 不是点云，只是逐点观测记录。",
            "L2 是相对飞机自身坐标系的局部点云。",
            "L3 把每个激光点与飞机 POS 轨迹按时间对应起来。",
            "最终结果是把“飞机位置”和“激光点相对飞机的偏移”合成。",
        ],
        size=13.2,
    )

    add_panel(slide, 6.95, 3.55, 5.65, 2.45, "本次阶段性结论")
    add_bullets(
        slide,
        7.22,
        4.02,
        5.1,
        1.65,
        [
            "00050-00150 中 95 个 Stage 7A 文件判为可展示候选成果。",
            "6 个坏时间文件单独排除，不混入几何判断。",
            "00144-00145 存在 0.12666 秒时间间隔，但目视未见明显几何跳变。",
        ],
        size=13.2,
    )
    add_note(slide, "汇报定位：解释处理链条和阶段性成果，不宣称已完成全部文件最终生产。")


def build_slide2(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_title(slide, "1. L1 原始数据：只有观测记录，还不是空间点云", "示例文件：L1_cap_00111_20260510190957.h5")
    add_picture_fit(slide, FIG_DIR / "stage7d_00111_L1_raw.png", 0.55, 1.25, 6.55, 5.55)

    add_panel(slide, 7.35, 1.25, 5.3, 5.55, "L1 中主要记录什么？")
    add_bullets(
        slide,
        7.68,
        1.75,
        4.75,
        1.45,
        [
            "每个激光点的时间：GNSS_SEC_CH1",
            "原始距离计数：Photon_CH1_DIST",
            "扫描镜位置/相位：Photon_CH1_CODER",
            "脉冲索引：PULSE_INDEX_CH1",
        ],
        size=12.8,
    )
    add_textbox(slide, 7.68, 3.55, 4.65, 0.32, "一句话解释", size=14, bold=True, color=NAVY)
    add_textbox(
        slide,
        7.68,
        3.93,
        4.65,
        1.15,
        "L1 告诉我们“什么时候测、测了多远、扫描镜在哪个角度”，但还不知道点在真实空间中的位置。",
        size=15,
        color=RGBColor(35, 35, 35),
    )
    add_textbox(slide, 7.68, 5.55, 4.65, 0.32, "图中看什么？", size=14, bold=True, color=NAVY)
    add_textbox(
        slide,
        7.68,
        5.92,
        4.65,
        0.55,
        "检查原始距离、扫描相位和脉冲索引随时间是否正常连续。",
        size=12.5,
        color=GRAY,
    )
    add_note(slide, "这一页只说明原始观测量，不涉及论文几何模型。")


def build_slide3(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_title(slide, "2. L1 到 L2：把距离和扫描角转换成飞机坐标系下的点", "这一步得到相对飞机的局部三维点云，还没有放到地图坐标里")
    add_picture_fit(slide, FIG_DIR / "stage7d_00111_L2_FRD.png", 0.55, 1.18, 6.35, 5.65)

    add_panel(slide, 7.15, 1.18, 5.55, 2.45, "处理逻辑")
    add_bullets(
        slide,
        7.48,
        1.64,
        5.0,
        1.35,
        [
            "先把原始距离计数换算成真实距离：range_m",
            "再把编码器位置换算成扫描角：scan_angle_deg",
            "最后计算激光点相对飞机的前后、左右、上下位置：frd_x_m / frd_y_m / frd_z_m",
        ],
        size=12.3,
    )

    add_panel(slide, 7.15, 3.9, 5.55, 2.35, "论文使用点：曹彬才 第 4.4.2 节")
    add_textbox(
        slide,
        7.48,
        4.38,
        4.95,
        1.28,
        "这里主要参考 LADM-II 扫描镜几何：用反射镜法线、入射光方向和反射光方向计算激光束出射方向。论文中涉及约 7.5° 反射镜倾角和 45° 坐标变换思想，对应公式 4-7 到 4-14。",
        size=12.3,
        color=RGBColor(35, 35, 35),
    )
    add_textbox(
        slide,
        7.48,
        5.78,
        4.95,
        0.32,
        "注意：没有声称完整照搬第 4.4 节所有直接地理定位公式。",
        size=11.2,
        bold=True,
        color=RED,
    )
    add_note(slide, "对老师可以说：这一页是把“测距 + 扫描镜角度”变成“相对飞机的三维位置”。")


def build_slide4(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_title(slide, "3. L2 到 L3，再到最终坐标：给每个点找到飞机位置和姿态", "最终坐标 = 飞机当时的位置 + 激光点相对飞机的偏移")
    add_picture_fit(slide, FIG_DIR / "stage7d_00111_L3_POS_NED.png", 0.55, 1.18, 6.35, 5.65)

    add_panel(slide, 7.15, 1.18, 5.55, 2.05, "L2 到 L3：POS 匹配")
    add_bullets(
        slide,
        7.48,
        1.62,
        4.95,
        1.12,
        [
            "按时间找到飞机当时的位置：pos_easting / pos_northing / pos_height",
            "插值得到飞机姿态：pos_roll / pos_pitch / pos_heading",
            "旋转后得到激光点相对飞机的东、北、下偏移：east_offset_m / north_offset_m / down_offset_m",
        ],
        size=11.7,
    )

    add_panel(slide, 7.15, 3.48, 5.55, 1.65, "最终坐标计算")
    add_textbox(slide, 7.48, 3.92, 4.95, 0.28, "最终东坐标 = 飞机东坐标 + 激光点东向偏移", size=12.2, color=RGBColor(35, 35, 35))
    add_textbox(slide, 7.48, 4.25, 4.95, 0.28, "最终北坐标 = 飞机北坐标 + 激光点北向偏移", size=12.2, color=RGBColor(35, 35, 35))
    add_textbox(slide, 7.48, 4.58, 4.95, 0.28, "最终高程 = 飞机高程 - 激光点向下偏移", size=12.2, color=RGBColor(35, 35, 35))

    add_panel(slide, 7.15, 5.42, 5.55, 1.28, "最终汇报结论")
    add_textbox(
        slide,
        7.48,
        5.82,
        4.95,
        0.55,
        "00111 可完整展示 L1、L2、L3 处理链；00050-00150 中 95 个 Stage 7A 文件已整理为可展示候选成果。",
        size=13,
        bold=True,
        color=NAVY,
    )
    add_note(slide, "最终点云图和 CloudCompare 检查已经在 Stage 7A/7D 材料包中给出；本 PPT 重点讲中间过程。")


def main() -> None:
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    prs = Presentation()
    prs.slide_width = Inches(13.333333)
    prs.slide_height = Inches(7.5)
    build_slide1(prs)
    build_slide2(prs)
    build_slide3(prs)
    build_slide4(prs)
    prs.save(PPTX_PATH)
    print(PPTX_PATH)


if __name__ == "__main__":
    main()
